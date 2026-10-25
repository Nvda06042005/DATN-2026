"""Kiểm thử API của dịch vụ AI (chạy trong thư mục ai/: python -m unittest test_server -v).

Không gọi Gemini/Hugging Face: khóa để trống nên dùng mô hình NLP cục bộ và câu hỏi dự phòng.
"""
import base64
import io
import os
import unittest
import zipfile
from datetime import datetime
from unittest.mock import patch

os.environ["GEMINI_API_KEY"] = ""
os.environ["HF_TOKEN"] = ""
os.environ["AI_SERVICE_TOKEN"] = ""

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

import excel_provisioning as xp  # noqa: E402
from config import settings  # noqa: E402
from server import app  # noqa: E402

client = TestClient(app)


class AiServiceTests(unittest.TestCase):
    def test_health(self):
        r = client.get("/health")
        self.assertEqual(r.status_code, 200)
        self.assertGreaterEqual(r.json()["units"], 36)

    def test_chat_answers_from_knowledge_base(self):
        r = client.post("/chat", json={"query": "Trần Khánh Dư là ai?", "session_id": "test-1"})
        self.assertEqual(r.status_code, 200, r.text)
        body = r.json()
        self.assertIn("buôn than", body["answer"])
        self.assertTrue(body["grounding_sources"])

    def test_chat_rejects_out_of_domain(self):
        r = client.post("/chat", json={"query": "Công thức tính diện tích hình tròn?", "session_id": "test-2"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["grounding_sources"], [])

    def test_bkt_update_is_stateless(self):
        first = client.post("/bkt/update", json={"state": {}, "unit_id": "KC1_OVERVIEW", "is_correct": True}).json()
        self.assertEqual(first["attempts"], 1)
        self.assertGreater(first["new_p_l"], first["old_p_l"])
        self.assertEqual(first["state"]["KC1_OVERVIEW"]["correct_count"], 1)
        second = client.post("/bkt/update", json={"state": first["state"], "unit_id": "KC1_OVERVIEW", "is_correct": False}).json()
        self.assertEqual(second["attempts"], 2)
        self.assertLess(second["new_p_l"], first["new_p_l"])
        # Dịch vụ không giữ trạng thái: gửi lại trạng thái rỗng thì bắt đầu từ đầu
        again = client.post("/bkt/update", json={"state": {}, "unit_id": "KC1_OVERVIEW", "is_correct": True}).json()
        self.assertEqual(again["attempts"], 1)

    def test_bkt_recommend(self):
        state = {"KC3_VAN_DON": {"p_l": 0.12, "last_updated_time": 0, "attempts": 3, "correct_count": 0}}
        recs = client.post("/bkt/recommend", json={"state": state, "top_k": 3}).json()["recommendations"]
        self.assertLessEqual(len(recs), 3)
        self.assertTrue(all("unit_id" in r and "topic" in r for r in recs))

    def test_quiz_generate_falls_back_without_llm(self):
        r = client.post("/quiz/generate", json={"topic": "Trận Bạch Đằng 1288", "num_questions": 2})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertGreater(r.json()["total_generated"], 0)

    def test_game_questions_skip_fixed_fallback(self):
        r = client.post("/questions/generate", json={"unit_id": "KC3_VAN_DON", "topic": "Trận Vân Đồn", "count": 2})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["questions"], [])

    def test_internal_token(self):
        with patch.object(settings, "AI_SERVICE_TOKEN", "bi-mat"):
            self.assertEqual(client.post("/bkt/recommend", json={"state": {}}).status_code, 401)
            ok = client.post("/bkt/recommend", json={"state": {}}, headers={"X-AI-Token": "bi-mat"})
            self.assertEqual(ok.status_code, 200)
            self.assertEqual(client.get("/health").status_code, 200)
            self.assertEqual(client.get("/excel/template").status_code, 401)


def _template() -> bytes:
    r = client.get("/excel/template")
    assert r.status_code == 200, r.text
    return r.content


def _fill(raw: bytes, rows: list[list], edit=None) -> bytes:
    """Điền dòng vào trang dữ liệu của file mẫu (openpyxl ghi được cả khi trang đang khóa trong Excel)"""
    wb = load_workbook(io.BytesIO(raw))
    ws = wb[xp.DATA_SHEET]
    for i, values in enumerate(rows, start=2):
        for j, v in enumerate(values, start=1):
            ws.cell(row=i, column=j, value=v)
    if edit:
        edit(wb)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _validate(raw: bytes, filename: str = "danh-sach.xlsx", mime: str = xp.XLSX_MIME) -> dict:
    r = client.post("/excel/validate", json={"file_b64": base64.b64encode(raw).decode(), "filename": filename, "mime": mime})
    assert r.status_code == 200, r.text
    return r.json()


ANH = ["HS001", "THCS Văn Lang", "7480", "Miền Nam", "TP. Hồ Chí Minh", "Nguyễn  Vũ Đức Anh", "06/04/2012", "2023-2024", "Anh@Example.edu.vn", "7", "hs"]


class ExcelTemplateTests(unittest.TestCase):
    def test_template_has_hidden_signed_sheet(self):
        wb = load_workbook(io.BytesIO(_template()))
        self.assertEqual(wb.sheetnames, [xp.DATA_SHEET, xp.GUIDE_SHEET, xp.SIGNATURE_SHEET])
        sig = wb[xp.SIGNATURE_SHEET]
        self.assertEqual(sig.sheet_state, "veryHidden")
        self.assertTrue(sig.protection.sheet)
        template_id = sig["B1"].value
        self.assertEqual(sig["B4"].value, xp.sign(template_id))
        self.assertEqual({p.name: p.value for p in wb.custom_doc_props}[xp.PROP_ID], template_id)
        self.assertEqual([c.value for c in wb[xp.DATA_SHEET][1]], xp.COLUMNS)
        self.assertTrue(wb[xp.DATA_SHEET].protection.sheet)
        self.assertFalse(wb[xp.DATA_SHEET]["B2"].protection.locked)
        self.assertTrue(wb.security.lockStructure)

    def test_rows_parsed_with_formula_credentials(self):
        raw = _fill(_template(), [
            ANH,
            ["GV001", "THCS Văn Lang", "07480", "Miền Nam", "TP. Hồ Chí Minh", "Phạm Văn Giáo", datetime(1985, 9, 2), "2023 – 2024", "giao@example.edu.vn", 7, "GV"],
            ["HS003", "THCS Văn Lang", "7480", "Miền Nam", "TP. Hồ Chí Minh", "Lê Minh Thiếu", None, "2023-2024", "thieu@example.edu.vn", "7", "HS"],
            ["HS004", "THCS Văn Lang", "7480", "", "TP. Hồ Chí Minh", "Đỗ Sai Mẫu", "31/02/2012", "2023-2025", "sai-email", "10", "PH"],
            [None, None, None, None, None, None, None, None, None, None, None],
        ])
        body = _validate(raw)
        self.assertTrue(body["ok"], body)
        rows = {r["row"]: r for r in body["rows"]}
        self.assertEqual(sorted(rows), [2, 3, 4, 5])
        anh = rows[2]
        self.assertTrue(anh["complete"])
        self.assertEqual(anh["values"]["full_name"], "Nguyễn Vũ Đức Anh")
        self.assertEqual(anh["values"]["email"], "anh@example.edu.vn")
        self.assertEqual(anh["values"]["birth_date"], "2012-04-06")
        self.assertEqual(anh["values"]["role"], "HS")
        self.assertEqual(anh["credentials"], {"username_prefix": "anh.2374807022", "password": "Nvda06042012VN@"})
        self.assertEqual(len(anh["embedding"]), 128)
        giao = rows[3]
        self.assertEqual(giao["values"]["school_code"], "7480")
        self.assertEqual(giao["values"]["school_year"], "2023-2024")
        self.assertEqual(giao["credentials"], {"username_prefix": "giao.2374807022", "password": "Pvg02091985VN@"})
        self.assertEqual(rows[4]["missing"], ["Ngày tháng năm sinh"])
        self.assertIsNone(rows[4]["credentials"])
        bad = rows[5]
        self.assertEqual(bad["missing"], ["Khu vực"])
        self.assertEqual({i["field"] for i in bad["invalid"]}, {"Ngày tháng năm sinh", "Năm học", "Email", "Khối", "Vai trò"})
        self.assertFalse(bad["complete"])

    def test_rejects_edited_structure(self):
        renamed = _validate(_fill(_template(), [ANH], lambda wb: wb[xp.DATA_SHEET].cell(row=1, column=3, value="Mã trường")))
        self.assertFalse(renamed["ok"])
        self.assertEqual(renamed["errors"], ['Cấu trúc cột đã bị sửa: cột 3 là "Mã trường" thay vì "Mã số trường".'])

        def swap(wb):
            ws = wb[xp.DATA_SHEET]
            ws["D1"], ws["E1"] = "Thành phố", "Khu vực"
        self.assertEqual(len(_validate(_fill(_template(), [ANH], swap))["errors"]), 2)

        extra = _validate(_fill(_template(), [ANH + ["ghi chú"]]))
        self.assertFalse(extra["ok"])
        self.assertIn("Có thêm cột ngoài mẫu", extra["errors"][0])

        def forged(wb):
            wb[xp.SIGNATURE_SHEET]["B4"] = "0" * 64
        self.assertIn("Chữ ký file mẫu không hợp lệ", _validate(_fill(_template(), [ANH], forged))["errors"][0])

        def props(wb):
            for p in wb.custom_doc_props:
                if p.name == xp.PROP_ID:
                    p.value = "khac"
        self.assertIn("Thuộc tính file mẫu đã bị thay đổi", _validate(_fill(_template(), [ANH], props))["errors"][0])

        def shown(wb):
            wb[xp.SIGNATURE_SHEET].sheet_state = "visible"
        self.assertFalse(_validate(_fill(_template(), [ANH], shown))["ok"])

    def test_rejects_self_made_and_non_xlsx(self):
        wb = Workbook()
        wb.active.title = xp.DATA_SHEET
        wb.active.append(xp.COLUMNS)
        wb.active.append(ANH)
        out = io.BytesIO()
        wb.save(out)
        own = _validate(out.getvalue())
        self.assertFalse(own["ok"])
        self.assertIn("không có chữ ký của hệ thống", own["errors"][0])
        self.assertIn("Chỉ nhận file .xlsx", _validate(b"a,b", "ds.csv", "text/csv")["errors"][0])
        self.assertIn("không phải file Excel", _validate(_template(), "ds.xlsx", "application/vnd.ms-excel")["errors"][0])
        self.assertEqual(_validate(b"khong phai zip")["errors"], ["Nội dung không phải file .xlsx chuẩn."])
        macro = io.BytesIO()
        with zipfile.ZipFile(macro, "w") as z:
            z.writestr("[Content_Types].xml", "<Types/>")
            z.writestr("xl/workbook.xml", "<workbook/>")
            z.writestr("xl/vbaProject.bin", b"macro")
        self.assertEqual(_validate(macro.getvalue())["errors"], ["File chứa macro – không nhận."])
        broken = client.post("/excel/validate", json={"file_b64": "@@@@", "filename": "a.xlsx", "mime": xp.XLSX_MIME}).json()
        self.assertEqual(broken["errors"], ["File gửi lên bị hỏng."])

    def test_smart_delta_matching(self):
        old = {"full_name": "Lê Minh Thiếu", "user_code": "", "email": "", "school_code": "7480", "birth_date": None}
        new = {"full_name": "Lê Minh Thiếu", "user_code": "HS003", "email": "thieu@example.edu.vn", "school_code": "7480", "birth_date": "2012-11-20"}
        emb_old, emb_new = xp.embed(xp.profile_text(old)), xp.embed(xp.profile_text(new))
        self.assertGreater(xp.cosine(emb_old, emb_new), 0.5)
        rows = [
            {"row": 5, "values": new, "embedding": emb_new, "candidates": [{"id": 11, "full_name": "Lê Minh Thiếu", "school_code": "7480", "row_data": old, "embedding": emb_old}]},
            {"row": 6, "values": {**new, "full_name": "Lê Minh Thiệu", "user_code": "HS777", "email": "x@example.edu.vn"}, "embedding": emb_new, "candidates": [{"id": 11, "full_name": "Lê Minh Thiếu", "row_data": old, "embedding": emb_old}]},
            {"row": 7, "values": {**new, "user_code": "HS008"}, "embedding": emb_new, "candidates": [{"id": 12, "user_code": "HS008", "row_data": {}, "embedding": None}]},
            {"row": 8, "values": new, "embedding": emb_new, "candidates": [{"id": 13, "full_name": "Lê Minh Thiếu", "row_data": {"birth_date": "2011-01-01"}, "embedding": emb_new}]},
        ]
        r = client.post("/provision/match", json={"rows": rows})
        self.assertEqual(r.status_code, 200, r.text)
        d = {x["row"]: x for x in r.json()["decisions"]}
        self.assertEqual(d[5]["log_id"], 11)
        self.assertTrue(d[5]["reason"].startswith("cùng họ tên và trường"))
        self.assertIsNone(d[6]["log_id"])
        self.assertEqual((d[7]["log_id"], d[7]["reason"]), (12, "trùng mã người dùng"))
        self.assertIsNone(d[8]["log_id"])


if __name__ == "__main__":
    unittest.main()
